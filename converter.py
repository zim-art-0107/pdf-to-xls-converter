import sys
import aspose.pdf as ap
import openpyxl

def remove_evaluation_warning(excel_path):
    try:
        wb = openpyxl.load_workbook(excel_path)
        for sheet in wb.worksheets:
            # Aspose usually puts its watermark in the first row or as an evaluation string
            rows_to_delete = []
            for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=5), start=1):
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and ("Evaluation" in cell.value or "Aspose" in cell.value):
                        rows_to_delete.append(row_idx)
                        break
            
            for row_idx in reversed(rows_to_delete):
                sheet.delete_rows(row_idx)
                
        wb.save(excel_path)
        wb.close()
    except Exception as e:
        print(f"Warning: Could not remove watermark: {str(e)}")

def convert_pdf_to_excel(pdf_path, output_path):
    try:
        # Load PDF document
        document = ap.Document(pdf_path)
        
        # Instantiate ExcelSave Options object. By default it maps to Excel cells.
        save_option = ap.ExcelSaveOptions()
        save_option.format = ap.ExcelSaveOptions.ExcelFormat.XLSX
        
        # Save the file into MS Excel format
        document.save(output_path, save_option)
        
        # Clean the watermark
        remove_evaluation_warning(output_path)
        
        print(f"Success: {output_path}")
        sys.exit(0)
    except Exception as e:
        print(f"Error: {str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Error: Missing arguments. Usage: python converter.py <input_pdf> <output_xlsx>")
        sys.exit(1)
        
    input_pdf = sys.argv[1]
    output_excel = sys.argv[2]
    
    convert_pdf_to_excel(input_pdf, output_excel)
